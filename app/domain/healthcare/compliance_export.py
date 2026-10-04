"""
Compliance report export.

build_compliance_workbook() returns an .xlsx (bytes) with:
  Summary              headline figures, all formulas over the other sheets
  Staffing             facility x role headcount and hours, with totals
  Hours reconciliation paid vs scheduled per person per pay period
  Credentials          every license, expiry from the board, days left
  Open issues          what each team needs to look at
  Audit trail          which source rows every figure was built from
  Source lineage       original value -> cleaned value for each record

Differences, statuses and totals are Excel formulas, so the workbook stays
correct if someone edits an hours value while investigating.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
INK = "1F1D2B"
MUTED = "6B6878"
HEADER_FILL = PatternFill("solid", fgColor="F1ECFE")
TOTAL_FILL = PatternFill("solid", fgColor="F7F6F3")
RED_FILL = PatternFill("solid", fgColor="FDE8E8")
AMBER_FILL = PatternFill("solid", fgColor="FEF3DC")
GREEN_FILL = PatternFill("solid", fgColor="E6F5EE")
GREY_FILL = PatternFill("solid", fgColor="F1F0EE")
THIN = Side(style="thin", color="E4E2DD")
BORDER = Border(bottom=THIN)

SOURCE_LABEL = {"hr": "HR roster", "payroll": "Payroll", "licenses": "Licensing board", "schedule": "Schedule"}
LINEAGE_FIELDS = [
    ("employee_name", "canonical_name"), ("first_name", "first_name"), ("last_name", "last_name"),
    ("name_on_license", "canonical_name"), ("facility", "facility"), ("facility_code", "facility"),
    ("job_title", "job_title"), ("job_code", "job_title"), ("role", "job_title"),
    ("license_number", "license_number"), ("license_type", "license_type"),
    ("license_expiration", "license_expiration"), ("expiration_date", "license_expiration"),
    ("hire_date", "hire_date"), ("phone", "phone"), ("hours_paid", "hours"),
    ("period_start", "period_start"), ("period_end", "period_end"),
]


def _date(iso: str):
    try:
        return datetime.strptime(iso, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def _font(bold=False, color=INK, size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def _header(ws, row: int, headers: list[str], widths: list[int]):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = _font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 22
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _body(ws, first_row: int, last_row: int, ncols: int):
    for r in range(first_row, last_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = _font()
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str) and len(cell.value) > 30)


def _status_rules(ws, col: str, first: int, last: int):
    rng = f"{col}{first}:{col}{last}"
    for text, fill in (("Expired", RED_FILL), ("Action needed", RED_FILL),
                       ("Paid more than scheduled", AMBER_FILL), ("Scheduled more than paid", AMBER_FILL),
                       ("Renew within 30 days", AMBER_FILL), ("Needs a look", AMBER_FILL),
                       ("Renew within 90 days", GREY_FILL), ("Heads-up", GREY_FILL),
                       ("Not compared", GREY_FILL),
                       ("Matches", GREEN_FILL), ("Current", GREEN_FILL)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}{first}="{text}"'], fill=fill))


def build_compliance_workbook(
    report: dict,
    issues: list[dict] | None = None,
    persons: list | None = None,
    raw_records: list | None = None,
    normalized_records: list | None = None,
    source_files: list[str] | None = None,
    organization: str = "Harborview Care Group",
) -> bytes:
    issues = issues or []
    persons = persons or []
    wb = Workbook()

    # ------------------------------------------------------------------
    # Staffing
    # ------------------------------------------------------------------
    st = wb.active
    st.title = "Staffing"
    headers = ["Facility", "Role", "Headcount", "Licensed staff", "Valid licenses",
               "Scheduled hours", "Paid hours", "Paid minus scheduled"]
    _header(st, 1, headers, [16, 30, 12, 14, 14, 16, 12, 20])
    matrix = report.get("staffing_matrix", [])
    r = 2
    for row in matrix:
        st.cell(r, 1, row["facility"]); st.cell(r, 2, row["role"])
        st.cell(r, 3, row["headcount"]); st.cell(r, 4, row["licensed"]); st.cell(r, 5, row["license_valid"])
        st.cell(r, 6, row["scheduled_hours"]); st.cell(r, 7, row["paid_hours"])
        st.cell(r, 8, f"=G{r}-F{r}")
        r += 1
    staff_last = max(r - 1, 2)
    _body(st, 2, staff_last, len(headers))
    st.cell(r, 1, "Total").font = _font(bold=True)
    for col in "CDEFGH":
        c = st[f"{col}{r}"]
        c.value = f"=SUM({col}2:{col}{staff_last})"
        c.font = _font(bold=True)
    for col in range(1, len(headers) + 1):
        st.cell(r, col).fill = TOTAL_FILL
    staff_total_row = r
    for rr in range(2, r + 1):
        for col in "FGH":
            st[f"{col}{rr}"].number_format = '#,##0.0;(#,##0.0);"-"'

    # ------------------------------------------------------------------
    # Hours reconciliation
    # ------------------------------------------------------------------
    hr_ws = wb.create_sheet("Hours reconciliation")
    headers = ["Name", "Employee ID", "Facility", "Role", "Period start", "Period end",
               "Scheduled hours", "Paid hours", "Paid minus scheduled", "Status",
               "Payroll rows", "Scheduled shifts", "Note"]
    _header(hr_ws, 1, headers, [20, 12, 12, 26, 12, 12, 11, 11, 13, 24, 14, 44, 34])
    rows = report.get("hours_reconciliation", [])
    r = 2
    for row in rows:
        hr_ws.cell(r, 1, row["name"]); hr_ws.cell(r, 2, row.get("employee_id", ""))
        hr_ws.cell(r, 3, row["facility"]); hr_ws.cell(r, 4, row["role"])
        hr_ws.cell(r, 5, _date(row["period_start"])); hr_ws.cell(r, 6, _date(row["period_end"]))
        hr_ws.cell(r, 7, row["scheduled_hours"] if row["scheduled_hours"] is not None else None)
        hr_ws.cell(r, 8, row["paid_hours"])
        hr_ws.cell(r, 9, f'=IF(G{r}="","",H{r}-G{r})')
        hr_ws.cell(r, 10, (f'=IF(G{r}="","Not compared",IF(ABS(I{r})<=0.5,"Matches",'
                           f'IF(I{r}>0,"Paid more than scheduled","Scheduled more than paid")))'))
        hr_ws.cell(r, 11, ", ".join(p for p in row.get("payroll_ids", []) if p))
        hr_ws.cell(r, 12, "; ".join(row.get("shifts", [])))
        hr_ws.cell(r, 13, row.get("note", ""))
        r += 1
    hours_last = max(r - 1, 2)
    _body(hr_ws, 2, hours_last, len(headers))
    for rr in range(2, hours_last + 1):
        hr_ws[f"E{rr}"].number_format = "yyyy-mm-dd"
        hr_ws[f"F{rr}"].number_format = "yyyy-mm-dd"
        for col in "GHI":
            hr_ws[f"{col}{rr}"].number_format = '0.0;-0.0;"0.0"'
    _status_rules(hr_ws, "J", 2, hours_last)

    # ------------------------------------------------------------------
    # Credentials
    # ------------------------------------------------------------------
    cr = wb.create_sheet("Credentials")
    headers = ["Name", "Facility", "Role", "License number", "License type",
               "Expires (licensing board)", "Expiry in HR", "Days left", "Status", "HR matches board"]
    _header(cr, 1, headers, [20, 12, 26, 16, 24, 16, 14, 10, 22, 14])
    # HR vs board values per person
    rec_index = {n.raw_record_id: n for n in (normalized_records or [])}
    hr_exp: dict[str, str] = {}
    board_exp: dict[str, str] = {}
    for p in persons:
        for rid in p.source_records:
            n = rec_index.get(rid)
            if not n:
                continue
            if n.source == "hr" and p.canonical_name not in hr_exp:
                hr_exp[p.canonical_name] = n.data.get("license_expiration", "")
            if n.source == "licenses" and p.canonical_name not in board_exp:
                board_exp[p.canonical_name] = n.data.get("license_expiration", "")
    creds = report.get("credentials", [])
    r = 2
    for c in creds:
        cr.cell(r, 1, c["name"]); cr.cell(r, 2, c["facility"]); cr.cell(r, 3, c["role"])
        cr.cell(r, 4, c["license_number"]); cr.cell(r, 5, c["license_type"])
        cr.cell(r, 6, _date(board_exp.get(c["name"]) or c["expires"]))
        cr.cell(r, 7, _date(hr_exp.get(c["name"], "")))
        cr.cell(r, 8, f'=IF(F{r}="","",F{r}-Summary!$B$6)')
        cr.cell(r, 9, (f'=IF(H{r}="","Date missing",IF(H{r}<0,"Expired",IF(H{r}<=30,"Renew within 30 days",'
                       f'IF(H{r}<=90,"Renew within 90 days","Current"))))'))
        cr.cell(r, 10, f'=IF(OR(F{r}="",G{r}=""),"Only one source",IF(F{r}=G{r},"Yes","No"))')
        r += 1
    cred_last = max(r - 1, 2)
    _body(cr, 2, cred_last, len(headers))
    for rr in range(2, cred_last + 1):
        cr[f"F{rr}"].number_format = "yyyy-mm-dd"
        cr[f"G{rr}"].number_format = "yyyy-mm-dd"
        cr[f"H{rr}"].number_format = "0"
    _status_rules(cr, "I", 2, cred_last)
    cr.conditional_formatting.add(f"J2:J{cred_last}", FormulaRule(formula=['J2="No"'], fill=AMBER_FILL))

    # ------------------------------------------------------------------
    # Open issues
    # ------------------------------------------------------------------
    iss = wb.create_sheet("Open issues")
    headers = ["Level", "Team", "Name", "Facility", "Issue", "Details", "Next step"]
    _header(iss, 1, headers, [16, 18, 20, 12, 40, 52, 56])
    r = 2
    for i in issues:
        iss.cell(r, 1, i["level_label"]); iss.cell(r, 2, ", ".join(i["teams"]))
        iss.cell(r, 3, i["name"]); iss.cell(r, 4, i["facility"])
        iss.cell(r, 5, i["title"]); iss.cell(r, 6, i["detail"]); iss.cell(r, 7, i["next_step"])
        r += 1
    issues_last = max(r - 1, 2)
    _body(iss, 2, issues_last, len(headers))
    _status_rules(iss, "A", 2, issues_last)

    # ------------------------------------------------------------------
    # Audit trail
    # ------------------------------------------------------------------
    au = wb.create_sheet("Audit trail")
    headers = ["Person", "Figure", "Pay period", "Value", "Built from"]
    _header(au, 1, headers, [20, 16, 22, 10, 80])
    r = 2
    for a in report.get("audit_trail", []):
        au.cell(r, 1, a["person"]); au.cell(r, 2, a["figure"]); au.cell(r, 3, a["period"])
        au.cell(r, 4, a["value"]); au.cell(r, 5, a["built_from"])
        r += 1
    _body(au, 2, max(r - 1, 2), len(headers))

    # ------------------------------------------------------------------
    # Source lineage
    # ------------------------------------------------------------------
    ln = wb.create_sheet("Source lineage")
    headers = ["Person", "System", "File", "Column", "Original value", "Cleaned value"]
    _header(ln, 1, headers, [20, 16, 22, 20, 28, 28])
    raw_index = {x.id: x for x in (raw_records or [])}
    r = 2
    for p in persons:
        for rid in p.source_records:
            raw = raw_index.get(rid)
            norm = rec_index.get(rid)
            if not raw or not norm or raw.source == "schedule":
                continue
            mapped = raw.data.get("mapped", {})
            colmap = raw.data.get("column_mapping", {})
            for raw_key, norm_key in LINEAGE_FIELDS:
                if raw_key not in mapped or mapped.get(raw_key) in ("", None):
                    continue
                ln.cell(r, 1, p.canonical_name)
                ln.cell(r, 2, SOURCE_LABEL.get(raw.source, raw.source))
                ln.cell(r, 3, raw.source_file)
                ln.cell(r, 4, colmap.get(raw_key, raw_key))
                ln.cell(r, 5, str(mapped.get(raw_key)))
                ln.cell(r, 6, str(norm.data.get(norm_key, "")))
                r += 1
    _body(ln, 2, max(r - 1, 2), len(headers))

    # ------------------------------------------------------------------
    # Summary (first sheet)
    # ------------------------------------------------------------------
    sm = wb.create_sheet("Summary", 0)
    sm.column_dimensions["A"].width = 44
    sm.column_dimensions["B"].width = 26
    sm.column_dimensions["C"].width = 60
    sm["A1"] = f"{organization} — Staffing and compliance report"
    sm["A1"].font = _font(bold=True, size=14)
    rp = report.get("report_period", {})
    sm["A3"] = "Report period"; sm["B3"] = rp.get("label", "")
    sm["A4"] = "Hours compared for this period"; sm["B4"] = "Yes" if rp.get("hours_compared") else "No"
    sm["A5"] = "Period start"; sm["B5"] = _date(rp.get("start", "")); sm["B5"].number_format = "yyyy-mm-dd"
    sm["A6"] = "Generated on"; sm["B6"] = _date(report.get("generated_on", "")) or date.today()
    sm["B6"].number_format = "yyyy-mm-dd"
    sm["A7"] = "Source files"; sm["B7"] = ", ".join(source_files or [])

    sm["A9"] = "Figure"; sm["B9"] = "Value"; sm["C9"] = "How it's calculated"
    for col in "ABC":
        sm[f"{col}9"].font = _font(bold=True); sm[f"{col}9"].fill = HEADER_FILL; sm[f"{col}9"].border = BORDER

    hrs = "'Hours reconciliation'"
    figures = [
        ("Staff on file", f"=Staffing!C{staff_total_row}", "Total headcount on the Staffing sheet"),
        ("Licensed staff", f"=Staffing!D{staff_total_row}", "People with a license number"),
        ("Valid licenses at period end", f"=Staffing!E{staff_total_row}", "Licensed staff whose license is valid on the period's last day"),
        ("Licenses expired today", f'=COUNTIF(Credentials!I2:I{cred_last},"Expired")', "Credentials sheet, status Expired"),
        ("Licenses to renew within 30 days", f'=COUNTIF(Credentials!I2:I{cred_last},"Renew within 30 days")', "Credentials sheet"),
        ("HR expiry differs from the board", f'=COUNTIF(Credentials!J2:J{cred_last},"No")', "Credentials sheet, HR matches board = No"),
        ("Scheduled hours (report period)", f"=SUMIFS({hrs}!G2:G{hours_last},{hrs}!E2:E{hours_last},$B$5)", "Hours reconciliation rows in the report period"),
        ("Paid hours (report period)", f"=SUMIFS({hrs}!H2:H{hours_last},{hrs}!E2:E{hours_last},$B$5)", "Hours reconciliation rows in the report period"),
        ("Paid minus scheduled", "=B17-B16", "Positive means more hours were paid than scheduled"),
        ("People whose hours don't match", (f'=COUNTIFS({hrs}!J2:J{hours_last},"Paid more than scheduled",{hrs}!E2:E{hours_last},$B$5)'
                                            f'+COUNTIFS({hrs}!J2:J{hours_last},"Scheduled more than paid",{hrs}!E2:E{hours_last},$B$5)'),
         "Difference larger than 0.5h in the report period"),
        ("Open issues: action needed", f'=COUNTIF(\'Open issues\'!A2:A{issues_last},"Action needed")', "Open issues sheet"),
        ("Open issues: needs a look", f'=COUNTIF(\'Open issues\'!A2:A{issues_last},"Needs a look")', "Open issues sheet"),
        ("Open issues: heads-up", f'=COUNTIF(\'Open issues\'!A2:A{issues_last},"Heads-up")', "Open issues sheet"),
    ]
    r = 10
    for label, formula, how in figures:
        sm.cell(r, 1, label); sm.cell(r, 2, formula); sm.cell(r, 3, how)
        sm.cell(r, 2).alignment = Alignment(horizontal="right")
        r += 1
    _body(sm, 10, r - 1, 3)
    for rr in (16, 17, 18):
        sm[f"B{rr}"].number_format = '#,##0.0;(#,##0.0);"-"'

    notes = report.get("notes", [])
    r += 1
    sm.cell(r, 1, "Notes").font = _font(bold=True)
    r += 1
    base_notes = [
        "The licensing board is treated as the source of truth for license numbers and expiry dates.",
        "Hours are compared only for pay periods the schedule fully covers. Other periods are listed as Not compared.",
    ]
    for n in base_notes + notes:
        c = sm.cell(r, 1, n)
        c.font = _font(color=MUTED)
        sm.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        sm.row_dimensions[r].height = 28
        r += 1
    for rr in range(3, 8):
        sm[f"A{rr}"].font = _font(color=MUTED)
        sm[f"B{rr}"].font = _font()
        sm[f"B{rr}"].alignment = Alignment(horizontal="left")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_hours_csv(report: dict) -> bytes:
    """Plain CSV of the hours reconciliation, for systems that can't read xlsx."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["name", "employee_id", "facility", "role", "period_start", "period_end",
                "scheduled_hours", "paid_hours", "paid_minus_scheduled", "status",
                "payroll_rows", "scheduled_shifts", "note"])
    for r in report.get("hours_reconciliation", []):
        w.writerow([
            r["name"], r.get("employee_id", ""), r["facility"], r["role"],
            r["period_start"], r["period_end"],
            "" if r["scheduled_hours"] is None else f"{r['scheduled_hours']:g}",
            f"{r['paid_hours']:g}",
            "" if r["difference"] is None else f"{r['difference']:g}",
            r["status"], " ".join(p for p in r.get("payroll_ids", []) if p),
            "; ".join(r.get("shifts", [])), r.get("note", ""),
        ])
    return buf.getvalue().encode("utf-8")
